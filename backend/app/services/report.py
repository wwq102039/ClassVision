"""报表服务：Excel 考勤报表导出。"""
from __future__ import annotations

import datetime as dt
import io

from sqlalchemy.orm import Session

from ..core import security
from ..core.exceptions import NoDataError
from ..models import Attendance, Course, Student


def export_excel(
    db: Session,
    class_name: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    course_id: int | None = None,
) -> bytes:
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font

    query = (
        db.query(Attendance, Student, Course)
        .join(Student, Attendance.student_id == Student.id)
        .join(Course, Attendance.course_id == Course.id)
    )
    if class_name:
        query = query.filter(Student.class_name == class_name)
    if course_id:
        query = query.filter(Attendance.course_id == course_id)
    if date_from:
        query = query.filter(Attendance.session_date >= dt.date.fromisoformat(date_from))
    if date_to:
        query = query.filter(Attendance.session_date <= dt.date.fromisoformat(date_to))
    rows = query.order_by(Attendance.sign_time.desc()).all()

    if not rows:
        raise NoDataError("暂无符合筛选条件的考勤数据")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "考勤报表"
    headers = ["学号", "姓名", "班级", "课程", "日期", "签到时间", "状态", "抓拍缩略图"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for att, stu, course in rows:
        ws.append(
            [
                stu.student_no,
                stu.name,
                stu.class_name,
                course.name,
                att.session_date.isoformat() if att.session_date else "",
                att.sign_time.strftime("%Y-%m-%d %H:%M:%S") if att.sign_time else "",
                att.status,
                "",
            ]
        )
        if att.photo_path:
            raw = security.read_encrypted_file(att.photo_path)
            if raw:
                try:
                    from PIL import Image

                    img = Image.open(io.BytesIO(raw)).convert("RGB")
                    img.thumbnail((80, 80))
                    buf = io.BytesIO()
                    img.save(buf, format="PNG")
                    buf.seek(0)
                    xl_img = XLImage(buf)
                    ws.add_image(xl_img, f"H{ws.max_row}")
                except Exception:  # noqa: BLE001
                    pass

    # 列宽
    widths = [14, 12, 14, 18, 12, 20, 10, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()